import os
import uuid
from typing import Dict, Any, List, Optional
from datetime import datetime, timezone
from src.eureka.universe.canonical_state import Evidence, ExtractedEvidence

class ParserResult:
    def __init__(self, extracted_evidence: Optional[ExtractedEvidence], error: Optional[str] = None, reason_code: Optional[str] = None):
        self.extracted_evidence = extracted_evidence
        self.error = error
        self.reason_code = reason_code

class EvidenceParser:
    parser_id = "BaseParser"
    parser_version = "1.0.0"

    def can_parse(self, evidence: Evidence) -> bool:
        return False
        
    def parse(self, evidence: Evidence) -> ParserResult:
        raise NotImplementedError()

class TextParser(EvidenceParser):
    parser_id = "TextParser"
    parser_version = "1.0.0"
    
    def can_parse(self, evidence: Evidence) -> bool:
        return evidence.extension.lower() == ".txt" or evidence.media_type == "text/plain"
        
    def parse(self, evidence: Evidence) -> ParserResult:
        try:
            with open(evidence.content_reference, "r", encoding="utf-8") as f:
                content = f.read()
            extracted = ExtractedEvidence(
                extracted_evidence_id=f"EXT-{uuid.uuid4().hex[:8].upper()}",
                evidence_id=evidence.evidence_id,
                content_type="text/plain",
                text_blocks=[content],
                source_locations=["line=1"],
                extraction_method="TXT_DETERMINISTIC",
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                extraction_timestamp=datetime.now(timezone.utc).isoformat()
            )
            return ParserResult(extracted_evidence=extracted)
        except Exception as e:
            return ParserResult(extracted_evidence=None, error=str(e), reason_code="DOCUMENT_PARSE_ERROR")

class MarkdownParser(EvidenceParser):
    parser_id = "MarkdownParser"
    parser_version = "1.0.0"
    
    def can_parse(self, evidence: Evidence) -> bool:
        return evidence.extension.lower() == ".md" or evidence.media_type == "text/markdown"
        
    def parse(self, evidence: Evidence) -> ParserResult:
        try:
            with open(evidence.content_reference, "r", encoding="utf-8") as f:
                content = f.read()
            # In a real implementation we would parse markdown blocks. For now, text blocks are paragraphs.
            blocks = [b.strip() for b in content.split("\n\n") if b.strip()]
            extracted = ExtractedEvidence(
                extracted_evidence_id=f"EXT-{uuid.uuid4().hex[:8].upper()}",
                evidence_id=evidence.evidence_id,
                content_type="text/markdown",
                text_blocks=blocks,
                source_locations=[f"block={i+1}" for i in range(len(blocks))],
                extraction_method="MD_DETERMINISTIC",
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                extraction_timestamp=datetime.now(timezone.utc).isoformat()
            )
            return ParserResult(extracted_evidence=extracted)
        except Exception as e:
            return ParserResult(extracted_evidence=None, error=str(e), reason_code="DOCUMENT_PARSE_ERROR")

class DocxParser(EvidenceParser):
    parser_id = "DocxParser"
    parser_version = "1.0.0"
    
    def can_parse(self, evidence: Evidence) -> bool:
        ext = evidence.extension.lower()
        mime = evidence.media_type
        return ext == ".docx" or "wordprocessingml.document" in mime
        
    def parse(self, evidence: Evidence) -> ParserResult:
        try:
            import docx
            doc = docx.Document(evidence.content_reference)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            tables = []
            for t_idx, table in enumerate(doc.tables):
                table_data = []
                for row in table.rows:
                    row_data = [cell.text for cell in row.cells]
                    table_data.append(row_data)
                if table_data:
                    tables.append({"table_index": t_idx, "data": table_data})
                    
            extracted = ExtractedEvidence(
                extracted_evidence_id=f"EXT-{uuid.uuid4().hex[:8].upper()}",
                evidence_id=evidence.evidence_id,
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                text_blocks=paragraphs,
                tables=tables,
                source_locations=[f"paragraph={i+1}" for i in range(len(paragraphs))],
                extraction_method="DOCX_DETERMINISTIC",
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                extraction_timestamp=datetime.now(timezone.utc).isoformat()
            )
            return ParserResult(extracted_evidence=extracted)
        except Exception as e:
            return ParserResult(extracted_evidence=None, error=str(e), reason_code="DOCUMENT_PARSE_ERROR")

class PdfParser(EvidenceParser):
    parser_id = "PdfParser"
    parser_version = "1.0.0"
    
    def can_parse(self, evidence: Evidence) -> bool:
        return evidence.extension.lower() == ".pdf" or evidence.media_type == "application/pdf"
        
    def parse(self, evidence: Evidence) -> ParserResult:
        try:
            import pypdf
            reader = pypdf.PdfReader(evidence.content_reference)
            pages_text = []
            source_locations = []
            has_scanned_pages = False
            has_text_pages = False
            
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text and text.strip():
                    pages_text.append(text.strip())
                    source_locations.append(f"page={i+1}")
                    has_text_pages = True
                else:
                    has_scanned_pages = True
                    
            if has_scanned_pages and not has_text_pages:
                # OCR Unavailable fallback
                return ParserResult(extracted_evidence=None, error="PDF is scanned/image-only, and OCR infrastructure is not installed.", reason_code="OCR_UNAVAILABLE")
                
            extracted = ExtractedEvidence(
                extracted_evidence_id=f"EXT-{uuid.uuid4().hex[:8].upper()}",
                evidence_id=evidence.evidence_id,
                content_type="application/pdf",
                text_blocks=pages_text,
                pages=pages_text,
                source_locations=source_locations,
                extraction_method="PDF_DETERMINISTIC",
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                extraction_timestamp=datetime.now(timezone.utc).isoformat(),
                warnings=["Some pages were scanned and OCR is unavailable"] if has_scanned_pages else []
            )
            return ParserResult(extracted_evidence=extracted)
        except Exception as e:
            return ParserResult(extracted_evidence=None, error=str(e), reason_code="DOCUMENT_PARSE_ERROR")

class CsvParser(EvidenceParser):
    parser_id = "CsvParser"
    parser_version = "1.0.0"
    
    def can_parse(self, evidence: Evidence) -> bool:
        return evidence.extension.lower() == ".csv" or evidence.media_type == "text/csv"
        
    def parse(self, evidence: Evidence) -> ParserResult:
        try:
            import pandas as pd
            df = pd.read_csv(evidence.content_reference)
            
            extracted = ExtractedEvidence(
                extracted_evidence_id=f"EXT-{uuid.uuid4().hex[:8].upper()}",
                evidence_id=evidence.evidence_id,
                content_type="text/csv",
                structured_data={
                    "columns": list(df.columns),
                    "row_count": len(df),
                    "column_count": len(df.columns),
                    "sample": df.head(5).to_dict(orient="records")
                },
                extraction_method="CSV_PANDAS",
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                extraction_timestamp=datetime.now(timezone.utc).isoformat()
            )
            return ParserResult(extracted_evidence=extracted)
        except Exception as e:
            return ParserResult(extracted_evidence=None, error=str(e), reason_code="DOCUMENT_PARSE_ERROR")

class XlsxParser(EvidenceParser):
    parser_id = "XlsxParser"
    parser_version = "1.0.0"
    
    def can_parse(self, evidence: Evidence) -> bool:
        ext = evidence.extension.lower()
        mime = evidence.media_type
        return ext == ".xlsx" or "spreadsheetml.sheet" in mime
        
    def parse(self, evidence: Evidence) -> ParserResult:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(evidence.content_reference, data_only=True)
            sheets = []
            source_locations = []
            
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                if ws.max_row > 0 and ws.max_column > 0:
                    sheets.append(sheet_name)
                    source_locations.append(f"sheet={sheet_name}")
                    
            if not sheets:
                raise ValueError("No non-empty sheets found in workbook")
                
            extracted = ExtractedEvidence(
                extracted_evidence_id=f"EXT-{uuid.uuid4().hex[:8].upper()}",
                evidence_id=evidence.evidence_id,
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                sheets=sheets,
                structured_data={"sheet_names": sheets},
                source_locations=source_locations,
                extraction_method="XLSX_OPENPYXL",
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                extraction_timestamp=datetime.now(timezone.utc).isoformat()
            )
            return ParserResult(extracted_evidence=extracted)
        except Exception as e:
            return ParserResult(extracted_evidence=None, error=str(e), reason_code="DOCUMENT_PARSE_ERROR")

class PptxParser(EvidenceParser):
    parser_id = "PptxParser"
    parser_version = "1.0.0"
    
    def can_parse(self, evidence: Evidence) -> bool:
        ext = evidence.extension.lower()
        mime = evidence.media_type
        return ext == ".pptx" or "presentationml.presentation" in mime
        
    def parse(self, evidence: Evidence) -> ParserResult:
        try:
            import pptx
            prs = pptx.Presentation(evidence.content_reference)
            slides = []
            source_locations = []
            
            for i, slide in enumerate(prs.slides):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        slide_text.append(shape.text.strip())
                if slide_text:
                    slides.append("\n".join(slide_text))
                    source_locations.append(f"slide={i+1}")
                    
            extracted = ExtractedEvidence(
                extracted_evidence_id=f"EXT-{uuid.uuid4().hex[:8].upper()}",
                evidence_id=evidence.evidence_id,
                content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                slides=slides,
                text_blocks=slides,
                source_locations=source_locations,
                extraction_method="PPTX_PYTHONPPTX",
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                extraction_timestamp=datetime.now(timezone.utc).isoformat()
            )
            return ParserResult(extracted_evidence=extracted)
        except Exception as e:
            return ParserResult(extracted_evidence=None, error=str(e), reason_code="DOCUMENT_PARSE_ERROR")

class ImageParser(EvidenceParser):
    parser_id = "ImageParser"
    parser_version = "1.0.0"
    
    def can_parse(self, evidence: Evidence) -> bool:
        ext = evidence.extension.lower()
        mime = evidence.media_type
        return ext in [".png", ".jpg", ".jpeg"] or mime.startswith("image/")
        
    def parse(self, evidence: Evidence) -> ParserResult:
        # Gracefully handle OCR unavailability without rejecting the valid evidence
        return ParserResult(
            extracted_evidence=None, 
            error="OCR capability is currently unavailable in this environment.", 
            reason_code="OCR_UNAVAILABLE"
        )

class EvidenceParserRegistry:
    def __init__(self):
        self.parsers: List[EvidenceParser] = [
            TextParser(),
            MarkdownParser(),
            DocxParser(),
            PdfParser(),
            CsvParser(),
            XlsxParser(),
            PptxParser(),
            ImageParser()
        ]
        
    def get_parser(self, evidence: Evidence) -> Optional[EvidenceParser]:
        for parser in self.parsers:
            if parser.can_parse(evidence):
                return parser
        return None
